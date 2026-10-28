"""
Lectura del Excel de control (CONTROL_ANS_SOX.xlsx).

* Detecta la fila de encabezados (la que contiene SERVIDOR y AGENDA) aunque haya títulos arriba.
* Rellena celdas combinadas con el valor de la celda ancla.
* Devuelve las filas con su número real de Excel para trazabilidad.
"""
from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from .agenda import Regla, interpretar_agenda, normalizar
from .servidores import Servidor, extraer_servidores

COLUMNAS_SALIDA = [
    "ANS", "SERVICIO", "SERVIDOR", "NOMBRE_POLITICA", "RUTAS", "TIPO_BACKUP", "AGENDA",
    "HORA_EJECUCION", "RETENCION (SEMANAS)", "TIPO_EJECUCION", "PLATAFORMA_BACKUP",
    "SERVER_BACKUP", "FECHA DE ACTUALIZACION", "OBSERVACIÒN", "REVISADO",
]


@dataclass
class FilaControl:
    fila_excel: int
    datos: dict
    servidor_texto: str
    agenda_texto: str
    servidores: list[Servidor] = field(default_factory=list)
    regla: Regla | None = None


def listar_hojas(ruta: str | Path) -> list[str]:
    wb = load_workbook(ruta, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _a_texto(v) -> str:
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d") if v.time() == dt.time(0) else v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, dt.date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, dt.time):
        return v.strftime("%H:%M")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def leer_control(ruta: str | Path, hoja: str | None = None) -> tuple[list[FilaControl], list[str], list[str]]:
    """Retorna (filas, encabezados, avisos)."""
    avisos: list[str] = []
    wb = load_workbook(ruta, data_only=True)
    try:
        if hoja and hoja in wb.sheetnames:
            ws = wb[hoja]
        else:
            if hoja:
                avisos.append(f"No existe la hoja '{hoja}'; se usó '{wb.sheetnames[0]}'")
            ws = wb[wb.sheetnames[0]]

        # valores de celdas combinadas
        combinadas: dict[tuple[int, int], object] = {}
        for rango in ws.merged_cells.ranges:
            ancla = ws.cell(rango.min_row, rango.min_col).value
            for r in range(rango.min_row, rango.max_row + 1):
                for c in range(rango.min_col, rango.max_col + 1):
                    combinadas[(r, c)] = ancla

        def valor(r: int, c: int):
            v = ws.cell(r, c).value
            if v is None and (r, c) in combinadas:
                v = combinadas[(r, c)]
            return v

        max_col = ws.max_column
        fila_enc = None
        for r in range(1, min(ws.max_row, 25) + 1):
            vals = [normalizar(_a_texto(valor(r, c))) for c in range(1, max_col + 1)]
            if "servidor" in vals and "agenda" in vals:
                fila_enc = r
                break
        if fila_enc is None:
            raise ValueError("No encontré una fila de encabezados con las columnas SERVIDOR y AGENDA")

        encabezados: list[str] = []
        indices: list[int] = []
        for c in range(1, max_col + 1):
            h = _a_texto(valor(fila_enc, c)).strip()
            if h:
                encabezados.append(h)
                indices.append(c)
        norm = [normalizar(h) for h in encabezados]
        i_srv = indices[norm.index("servidor")]
        i_age = indices[norm.index("agenda")]

        filas: list[FilaControl] = []
        vacias_seguidas = 0
        for r in range(fila_enc + 1, ws.max_row + 1):
            datos = {h: _a_texto(valor(r, c)) for h, c in zip(encabezados, indices)}
            if not any(v.strip() for v in datos.values()):
                vacias_seguidas += 1
                if vacias_seguidas > 50:
                    break
                continue
            vacias_seguidas = 0
            srv = _a_texto(valor(r, i_srv))
            age = _a_texto(valor(r, i_age))
            filas.append(FilaControl(
                fila_excel=r, datos=datos, servidor_texto=srv, agenda_texto=age,
                servidores=extraer_servidores(srv), regla=interpretar_agenda(age),
            ))
        sin_srv = sum(1 for f in filas if not f.servidores)
        if sin_srv:
            avisos.append(f"{sin_srv} filas no tienen un servidor reconocible")
        return filas, encabezados, avisos
    finally:
        wb.close()
