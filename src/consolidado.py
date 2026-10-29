"""
Consolidación de resultados y escritura del Excel de salida.

Estado por fila del Excel de control (decisión de Python, no de Copilot):
  SI                   -> todas las fechas pedidas (de todos sus servidores) tienen evidencia exacta
  PARCIAL              -> alguna coincide exacta y otras no
  CERCANA              -> ninguna exacta, pero hay fechas cercanas
  NO COINCIDEN FECHAS  -> no hay evidencia
  PENDIENTE / SIN RESPUESTA / SIN SERVIDOR / SIN FECHAS EN RANGO -> casos operativos
"""
from __future__ import annotations

import os
import time
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .agenda import normalizar
from .planificador import Plan

ESTADOS_FILA = ["SI", "PARCIAL", "CERCANA", "NO COINCIDEN FECHAS", "SIN RESPUESTA", "PENDIENTE",
                "SIN SERVIDOR", "SIN FECHAS EN RANGO"]
COLORES = {
    "SI": "C6EFCE", "PARCIAL": "FFEB9C", "CERCANA": "FCE4D6", "NO COINCIDEN FECHAS": "FFC7CE",
    "SIN RESPUESTA": "D9D9D9", "PENDIENTE": "EDEDED", "SIN SERVIDOR": "D9D9D9", "SIN FECHAS EN RANGO": "D9D9D9",
    "EXACTA": "C6EFCE", "NO_ENCONTRADA": "FFC7CE",
}
ROJO = "DA291C"


def _f(iso: str) -> str:
    return date.fromisoformat(iso).strftime("%d/%m/%Y") if iso else ""


def _dato(datos: dict, *nombres: str) -> str:
    norm = {normalizar(k): v for k, v in datos.items()}
    for n in nombres:
        v = norm.get(normalizar(n))
        if v:
            return v
    return ""


def evaluar_filas(plan: Plan, resultados: dict, estado_unidades: dict) -> list[dict]:
    """resultados: {uid: {fecha_iso: dict(ResultadoFecha)}}; estado_unidades: {uid: {"estado":..}}"""
    salida = []
    for fp in plan.filas:
        servidores = fp.servidores
        unidades_fila = plan.unidades_de_fila(fp)
        # grupos de la fila: 1 grupo (criterio "cualquiera") o 1 por servidor (criterio "cada_uno")
        grupos: dict[str, list] = {}
        for u in unidades_fila:
            grupos.setdefault(u.clave, []).append(u)
        varios = len(grupos) > 1
        pares = []            # (grupo, fecha_pedida, resultado|None, uid)
        pendiente = sin_resp = False
        for us in grupos.values():
            etiqueta_g = us[0].etiqueta
            if fp.origen == "Python":
                for iso in fp.fechas:
                    uid = next((u.id for u in us if iso in u.fechas), None)
                    res = resultados.get(uid, {}).get(iso) if uid else None
                    pares.append((etiqueta_g, iso, res, uid))
                    if res is None and uid:
                        st = estado_unidades.get(uid, {}).get("estado", "pendiente")
                        pendiente |= st == "pendiente"
                        sin_resp |= st == "sin_respuesta"
            else:
                for u in us:
                    rs = sorted(resultados.get(u.id, {}).values(), key=lambda r: r["fecha_pedida"])
                    for r in rs:
                        pares.append((etiqueta_g, r["fecha_pedida"], r, u.id))
                    if not rs:
                        st = estado_unidades.get(u.id, {}).get("estado", "pendiente")
                        pendiente |= st == "pendiente"
                        sin_resp |= st == "sin_respuesta"

        def etiqueta(srv, txt):
            return f"{srv} {txt}" if varios else txt

        exactas = [p for p in pares if p[2] and p[2]["estado"] == "EXACTA"]
        cercanas = [p for p in pares if p[2] and p[2]["estado"] == "CERCANA"]
        nulas = [p for p in pares if p[2] and p[2]["estado"] == "NO_ENCONTRADA"]
        sin_dato = [p for p in pares if not p[2]]
        respondidas = len(exactas) + len(cercanas) + len(nulas)
        obs = []

        if not servidores:
            estado = "SIN SERVIDOR"
        elif fp.origen == "Python" and not fp.fechas:
            estado = "SIN FECHAS EN RANGO"
        elif respondidas == 0:
            estado = "SIN RESPUESTA" if (sin_resp and not pendiente) else "PENDIENTE"
        else:
            if sin_dato:
                obs.append(f"{len(sin_dato)} fecha(s) sin respuesta de Copilot")
            total = respondidas + len(sin_dato)
            if len(exactas) == total:
                estado = "SI"
            elif exactas:
                estado = "PARCIAL"
            elif cercanas:
                estado = "CERCANA"
            elif sin_dato and (pendiente):
                estado = "PENDIENTE"
            else:
                estado = "NO COINCIDEN FECHAS"

        txt_exactas = ", ".join(etiqueta(p[0], _f(p[1])) for p in exactas)
        txt_cercanas = ", ".join(etiqueta(p[0], f"{_f(p[1])} → {_f(p[2]['fecha_encontrada'])}") for p in cercanas)
        txt_nulas = ", ".join(etiqueta(p[0], _f(p[1])) for p in nulas)
        if estado == "SI":
            resultado = "estado: SI"
        elif estado == "PARCIAL":
            partes = [f"coinciden: {txt_exactas}"]
            if txt_cercanas:
                partes.append(f"más cercanas: {txt_cercanas}")
            if txt_nulas:
                partes.append(f"sin evidencia: {txt_nulas}")
            resultado = "estado: PARCIAL | " + " | ".join(partes)
        elif estado == "CERCANA":
            resultado = f"estado: CERCANA | más cercanas: {txt_cercanas}"
            if txt_nulas:
                resultado += f" | sin evidencia: {txt_nulas}"
        else:
            resultado = f"estado: {estado}"

        evidencias = list(dict.fromkeys(p[2]["evidencia"] for p in exactas + cercanas if p[2].get("evidencia")))
        for p in pares:
            if p[2] and p[2].get("nota"):
                obs.append(p[2]["nota"])
        if fp.nota:
            obs.append(fp.nota)
        if fp.regla_nota and fp.origen == "Copilot":
            obs.append(fp.regla_nota)

        if fp.origen == "Python":
            solicitadas = ", ".join(_f(x) for x in fp.fechas)
        else:
            elegidas = sorted({p[1] for p in pares})
            solicitadas = (f"(Copilot elige {fp.n}) " + ", ".join(_f(x) for x in elegidas)).strip()

        salida.append({
            "FILA_EXCEL": fp.fila_excel,
            "ANS": _dato(fp.datos, "ANS"),
            "SERVICIO": _dato(fp.datos, "SERVICIO"),
            "SERVIDOR": fp.servidor_texto,
            "SERVIDORES_BUSCADOS": " / ".join(s["nombre"] for s in servidores),
            "CRITERIO": ("Basta con cualquiera" if plan.parametros.get("criterio", "cualquiera") == "cualquiera"
                         else "Cada servidor por separado") if len(servidores) > 1 else "",
            "NOMBRE_POLITICA": _dato(fp.datos, "NOMBRE_POLITICA", "NOMBRE POLITICA"),
            "TIPO_BACKUP": _dato(fp.datos, "TIPO_BACKUP", "TIPO BACKUP"),
            "AGENDA": " ".join(fp.agenda_texto.split()),
            "REGLA_INTERPRETADA": fp.regla,
            "ORIGEN_FECHAS": fp.origen,
            "N": fp.n,
            "FECHAS_SOLICITADAS": solicitadas,
            "FECHAS_COINCIDEN": txt_exactas,
            "FECHAS_CERCANAS": txt_cercanas,
            "FECHAS_SIN_EVIDENCIA": txt_nulas,
            "ESTADO": estado,
            "RESULTADO": resultado,
            "EVIDENCIA": " | ".join(evidencias)[:1000],
            "OBSERVACIONES": "; ".join(dict.fromkeys(obs))[:1000],
            "_pares": pares,
        })
    return salida


def contar_estados(filas_eval: list[dict]) -> Counter:
    return Counter(f["ESTADO"] for f in filas_eval)


# --------------------------------------------------------------------------------------
def _hoja_tabla(ws, columnas: list[str], filas: list[list], anchos: dict | None = None, col_estado: str | None = None):
    enc_fill = PatternFill("solid", fgColor=ROJO)
    enc_font = Font(bold=True, color="FFFFFF")
    borde = Border(bottom=Side(style="thin", color="DDDDDD"))
    ws.append(columnas)
    for c in range(1, len(columnas) + 1):
        cel = ws.cell(1, c)
        cel.fill, cel.font = enc_fill, enc_font
        cel.alignment = Alignment(vertical="center", wrap_text=True)
    for fila in filas:
        ws.append(fila)
    ws.freeze_panes = "B2"
    if filas:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{len(filas) + 1}"
    ws.row_dimensions[1].height = 30
    for i, col in enumerate(columnas, start=1):
        ancho = (anchos or {}).get(col)
        if ancho is None:
            largo = max([len(str(col))] + [len(str(f[i - 1])) for f in filas[:300] if f[i - 1] is not None])
            ancho = min(max(10, largo + 2), 60)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    if col_estado and col_estado in columnas:
        ci = columnas.index(col_estado) + 1
        for r in range(2, len(filas) + 2):
            cel = ws.cell(r, ci)
            color = COLORES.get(str(cel.value))
            if color:
                cel.fill = PatternFill("solid", fgColor=color)
                cel.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2):
        for cel in row:
            cel.alignment = Alignment(vertical="top", wrap_text=False)
            cel.border = borde


def escribir_excel(ruta: Path, plan: Plan, resultados: dict, estado_unidades: dict, lotes: list[dict],
                   info: dict) -> Path:
    filas_eval = evaluar_filas(plan, resultados, estado_unidades)
    wb = Workbook()

    # --- Resumen ---
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = "Auditoría de backups con Copilot"
    ws["A1"].font = Font(bold=True, size=16, color=ROJO)
    ws["A2"] = f"Generado: {datetime.now():%Y-%m-%d %H:%M}"
    ws["A2"].font = Font(color="666666")
    par = plan.parametros
    filas_info = [
        ("Rango de fechas", f"{_f(par['inicio'])} a {_f(par['fin'])}"),
        ("Semilla (reproducibilidad)", par.get("semilla", "")),
        ("Modo de fechas", "Python calcula; Copilot solo las agendas no reconocidas" if par.get("modo") == "hibrido"
         else "Copilot asigna todas las fechas"),
        ("Criterio con varios servidores", "Basta con que aparezca cualquiera"
         if par.get("criterio", "cualquiera") == "cualquiera" else "Cada servidor por separado"),
        ("Carpeta consultada", info.get("carpeta", "")),
        ("Ventana de cercanía (días)", info.get("ventana", "")),
        ("Estado de la ejecución", info.get("estado", "")),
        ("Lotes enviados a Copilot", len(lotes)),
    ]
    r = 4
    for k, v in filas_info:
        ws.cell(r, 1, k).font = Font(bold=True)
        ws.cell(r, 2, v)
        r += 1
    r += 1
    ws.cell(r, 1, "Estado").font = Font(bold=True, color="FFFFFF")
    ws.cell(r, 2, "Filas").font = Font(bold=True, color="FFFFFF")
    ws.cell(r, 1).fill = ws.cell(r, 2).fill = PatternFill("solid", fgColor=ROJO)
    conteo = contar_estados(filas_eval)
    for est in ESTADOS_FILA:
        if conteo.get(est):
            r += 1
            ws.cell(r, 1, est).fill = PatternFill("solid", fgColor=COLORES.get(est, "FFFFFF"))
            ws.cell(r, 2, conteo[est])
    r += 1
    ws.cell(r, 1, "TOTAL").font = Font(bold=True)
    ws.cell(r, 2, len(filas_eval)).font = Font(bold=True)
    r += 2
    ws.cell(r, 1, "Repeticiones por agenda").font = Font(bold=True)
    for k, v in sorted(par.get("repeticiones", {}).items()):
        r += 1
        ws.cell(r, 1, k)
        ws.cell(r, 2, v)
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 70

    # --- Consolidado ---
    cols = [k for k in (filas_eval[0].keys() if filas_eval else
                        ["FILA_EXCEL", "SERVIDOR", "AGENDA", "FECHAS_SOLICITADAS", "ESTADO", "RESULTADO"])
            if not k.startswith("_")]
    ws2 = wb.create_sheet("Consolidado")
    _hoja_tabla(ws2, cols, [[f[c] for c in cols] for f in filas_eval],
                anchos={"SERVIDOR": 34, "SERVIDORES_BUSCADOS": 34, "AGENDA": 24, "REGLA_INTERPRETADA": 30,
                        "FECHAS_SOLICITADAS": 40, "FECHAS_COINCIDEN": 36, "FECHAS_CERCANAS": 40,
                        "FECHAS_SIN_EVIDENCIA": 36, "RESULTADO": 70, "EVIDENCIA": 50, "OBSERVACIONES": 60,
                        "ESTADO": 22},
                col_estado="ESTADO")

    # --- Detalle por fecha ---
    det_cols = ["FILA_EXCEL", "SERVIDORES", "FECHA_PEDIDA", "ESTADO", "ESTADO_COPILOT", "FECHA_ENCONTRADA",
                "DIFERENCIA_DIAS", "EVIDENCIA", "NOTA_VALIDACION", "LOTE"]
    det = []
    for f in filas_eval:
        for srv, iso, res, uid in f["_pares"]:
            if res:
                det.append([f["FILA_EXCEL"], srv, _f(iso), res["estado"], res["estado_copilot"],
                            _f(res["fecha_encontrada"]), res["dif_dias"], res["evidencia"], res["nota"], res["lote"]])
            else:
                st = estado_unidades.get(uid, {}).get("estado", "pendiente") if uid else "pendiente"
                det.append([f["FILA_EXCEL"], srv, _f(iso), "SIN RESPUESTA" if st == "sin_respuesta" else "PENDIENTE",
                            "", "", None, "", "", None])
    ws3 = wb.create_sheet("Detalle por fecha")
    _hoja_tabla(ws3, det_cols, det, anchos={"EVIDENCIA": 50, "NOTA_VALIDACION": 50, "SERVIDORES": 34},
                col_estado="ESTADO")

    # --- Lotes (trazabilidad: prompt y respuesta cruda) ---
    lot_cols = ["LOTE", "PESTAÑA", "CODIGO", "ENVIADO", "DURACION_S", "ITEMS", "LINEAS_LEIDAS", "RESULTADO",
                "AVISOS", "PROMPT", "RESPUESTA"]
    ws4 = wb.create_sheet("Lotes Copilot")
    _hoja_tabla(ws4, lot_cols, [[l.get("n"), l.get("pestana"), l.get("codigo"), l.get("enviado"), l.get("duracion"),
                                 l.get("items"), l.get("lineas"), l.get("resultado"), "; ".join(l.get("avisos", []))[:2000],
                                 (l.get("prompt") or "")[:32000], (l.get("respuesta") or "")[:32000]] for l in lotes],
                anchos={"PROMPT": 60, "RESPUESTA": 80, "AVISOS": 50})

    ruta = Path(ruta)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    tmp = ruta.with_name(f"~tmp_{os.getpid()}_{ruta.name}")
    wb.save(tmp)
    try:
        os.replace(tmp, ruta)
        return ruta
    except PermissionError:
        # el Excel está abierto: dejar una copia con otro nombre
        alterna = ruta.with_name(f"{ruta.stem}_copia_{time.strftime('%H%M%S')}{ruta.suffix}")
        os.replace(tmp, alterna)
        return alterna
